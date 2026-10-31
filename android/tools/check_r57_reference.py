#!/usr/bin/env python3
"""Independently inspect R57 DOCX outputs with python-docx plus ZIP/XML; no rendering claim."""
import argparse
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]
WORD = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
REL = '{http://schemas.openxmlformats.org/package/2006/relationships}'


def inspect(directory):
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError('python-docx==1.2.0 is required as a separate host-only reader') from exc
    directory = Path(directory)
    count = 0

    def require(condition, message):
        nonlocal count
        count += 1
        if not condition:
            raise ValueError(message)

    structured = Document(directory / 'structured.docx')
    paragraphs = structured.paragraphs
    expected = ['Heading & <title>', 'Second', 'item bold and italic', 'quote code',
                '  plain\t🙂  ', '', 'last', '']
    require([p.text for p in paragraphs] == expected, 'structured text/line boundaries differ')
    require([p.style.name for p in paragraphs] == ['Heading 1', 'Heading 2', 'List Paragraph', 'Quote',
                                                   'Normal', 'Normal', 'Normal', 'Normal'], 'paragraph styles differ')
    runs = paragraphs[2].runs
    require([r.text for r in runs] == ['item ', 'bold', ' and ', 'italic'], 'inline run boundaries differ')
    require(runs[1].bold is True and runs[1].italic is not True, 'bold missing or incorrectly italic')
    require(runs[3].italic is True and runs[3].bold is not True, 'italic missing or incorrectly bold')
    require([r.text for r in paragraphs[3].runs] == ['quote ', 'code'], 'quote run boundaries differ')
    require(paragraphs[3].runs[1].font.name == 'monospace', 'monospace run missing')
    num_pr = paragraphs[2]._p.pPr.numPr
    require(num_pr is not None and num_pr.numId.val == 1 and num_pr.ilvl.val == 0, 'bullet numbering reference differs')
    for level, pt in enumerate((16, 14, 13, 12, 11, 10), 1):
        style = structured.styles[f'Heading {level}']
        require(style.font.bold is True and style.font.size is not None and style.font.size.pt == pt,
                f'Heading {level} definition missing bold/size')
        require(style.base_style is not None and style.base_style.name == 'Normal', f'Heading {level} base missing')
    require(structured.styles['Quote'].font.italic is True, 'quote style missing italic')
    with zipfile.ZipFile(directory / 'structured.docx') as package:
        names = package.namelist()
        require(len(names) == len(set(names)) and all('vbaProject' not in n for n in names), 'duplicate or macro part')
        rels = ET.fromstring(package.read('word/_rels/document.xml.rels'))
        pairs = {(r.attrib.get('Type', '').rsplit('/', 1)[-1], r.attrib.get('Target')) for r in rels.findall(REL + 'Relationship')}
        require(('styles', 'styles.xml') in pairs and ('numbering', 'numbering.xml') in pairs, 'styles/numbering relationship missing')
        numbering = ET.fromstring(package.read('word/numbering.xml'))
        nums = {n.attrib[WORD + 'numId']: n.find(WORD + 'abstractNumId').attrib[WORD + 'val'] for n in numbering.findall(WORD + 'num')}
        abstracts = {n.attrib[WORD + 'abstractNumId']: n for n in numbering.findall(WORD + 'abstractNum')}
        bullet = abstracts[nums['1']].find(WORD + 'lvl')
        require(bullet is not None and bullet.find(WORD + 'numFmt').attrib[WORD + 'val'] == 'bullet', 'numbering is not a bullet')
        styles = ET.fromstring(package.read('word/styles.xml'))
        headings = {n.attrib[WORD + 'styleId']: n for n in styles.findall(WORD + 'style')}
        for level in range(1, 7):
            outline = headings[f'Heading{level}'].find(f'{WORD}pPr/{WORD}outlineLvl')
            require(outline is not None and outline.attrib[WORD + 'val'] == str(level - 1), f'Heading {level} outline level differs')
    literal_expected = ['#tag', '####### not a heading', '# ', '- ', '> ', '| a | b |', '[link](https://example.invalid)',
                        '**unmatched', '*one **nested** end*', '***ambiguous***', '\\*escaped*', '`code **nested**`', 'x *not closed']
    literal = Document(directory / 'literal.docx')
    require([p.text for p in literal.paragraphs] == literal_expected, 'ambiguous/unsupported input was altered')
    require(all(p.style.name == 'Normal' for p in literal.paragraphs), 'literal lines became styled blocks')
    plain = Document(directory / 'plain.docx')
    require([p.text for p in plain.paragraphs] == ['# Heading & <title>', '## Second', '- item **bold** and *italic*',
                                                   '> quote `code`', '  plain\t🙂  ', '', 'last', ''], 'plain export changed text')
    require(all(p.style.name == 'Normal' for p in plain.paragraphs), 'plain export acquired block formatting')
    return count


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--directory', type=Path, default=ROOT / 'testing/build/r57/artifacts')
    args = parser.parse_args()
    count = inspect(args.directory)
    print(f'R57 independent python-docx and ZIP/XML checks: {count}; no Word/LibreOffice rendering or Android claim')


if __name__ == '__main__':
    main()
