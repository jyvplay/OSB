#!/usr/bin/env python3
"""Independent finite ROUND/OOXML checks. Neither reference engine is Excel."""
import argparse
from decimal import Decimal, ROUND_HALF_UP, localcontext
import math
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / 'testing/build/r56/artifacts'
SCALES = (-14, -13, -3, -2, -1, 0, 1, 2, 3, 4, 12, 13, 14, 20)
EDGES = {
    '=ROUND(1.005,2)': 1.01, '=ROUND(-1.005,2)': -1.01,
    '=ROUND(0.00000000000015,13)': 2e-13, '=ROUND(15000000000000,-13)': 2e13,
    '=ROUND(-15000000000000,-13)': -2e13, '=ROUND(1e308,2)': 1e308,
    '=ROUND(1e308,-308)': 1e308, '=ROUND(1e308,-309)': 0.0,
    '=ROUND(1e-320,320)': 1e-320, '=ROUND(1e-320,319)': 0.0,
    '=ROUND(1e-320,1e100)': 1e-320, '=ROUND(4.9e-324,323)': 0.0,
    '=ROUND(4.9e-324,324)': 5e-324, '=ROUND(4.9e-324,325)': 5e-324,
    '=ROUND(-4.9e-324,324)': -5e-324,
    '=ROUND(1.7e308,399)': 1.7e308, '=ROUND(1.7e308,400)': 1.7e308,
    '=ROUND(1.7e308,-399)': 0.0, '=ROUND(1.7e308,-400)': 0.0,
    '=ROUND(1.7976931348623157e308,399)': 1.7976931348623157e308,
    '=ROUND(1.7976931348623157e308,400)': 1.7976931348623157e308,
    '=ROUND(1.7976931348623157e308,-309)': 0.0,
    '=ROUND(1e308,-1e100)': 0.0, '=ROUND(2.5,0)': 3.0,
    '=ROUND(-2.5,0)': -3.0, '=ROUND(1.2345,1.9)': 1.2,
    '=ROUND(15.5,-1.9)': 20.0, '=ROUND(2.15,1)': 2.2,
    '=ROUND(-1.475,2)': -1.48, '=ROUND(626.3,-3)': 1000.0,
    '=IFERROR(ROUND(1.7e308,-308),42)': 42.0, '=ROUND(1.7e308,-308)': '#NUM!',
}
NONFINITE_REFERENCE = {
    '=ROUND(1.7e308,-308)', '=IFERROR(ROUND(1.7e308,-308),42)',
    '=ROUND(1.7976931348623157e308,399)', '=ROUND(1.7976931348623157e308,400)',
}


def same_number(actual, expected):
    """Typed, finite, exact-binary64 equality; zero never hides a subnormal."""
    if type(actual) not in (int, float) or type(expected) not in (int, float):
        return False
    try:
        return math.isfinite(actual) and math.isfinite(expected) and float(actual) == float(expected)
    except OverflowError:
        return False


def expected_keys():
    keys = set()
    for i in range(-200, 201):
        magnitude = abs(i * 5)
        literal = ('-' if i < 0 else '') + f'{magnitude // 1000}.{magnitude % 1000:03d}'
        keys.update((literal, scale) for scale in SCALES)
    return keys


def compare_rows(lines, keys, reference=None):
    seen = set()
    decimal_matches = reference_matches = 0
    for line in lines:
        fields = line.rstrip('\n').split('\t')
        if len(fields) != 3:
            raise ValueError('malformed corpus row')
        literal, raw_scale, raw_actual = fields
        try:
            scale, actual = int(raw_scale), float(raw_actual)
        except ValueError as exc:
            raise ValueError('nonnumeric corpus observation') from exc
        key = literal, scale
        if key not in keys or key in seen:
            raise ValueError(f'unexpected or duplicate corpus input: {key}')
        seen.add(key)
        with localcontext() as context:
            context.prec = 100
            decimal_result = Decimal(literal).quantize(Decimal(1).scaleb(-scale), rounding=ROUND_HALF_UP)
        expected = float(decimal_result)
        if not same_number(actual, expected):
            raise ValueError(f'Python decimal mismatch: {key}: got {actual}, expected {expected}')
        decimal_matches += 1
        if reference is not None:
            value = reference(f'=ROUND({literal},{scale})')
            if not same_number(actual, value):
                raise ValueError(f'formualizer mismatch: {key}: got {actual}, reference {value!r}')
            reference_matches += 1
    if seen != keys:
        raise ValueError(f'corpus incomplete: got {len(seen)}, expected {len(keys)}')
    return decimal_matches, reference_matches


def compare_edges(lines, expected, reference=None, nonfinite=frozenset()):
    seen = set()
    matches = divergences = 0
    for line in lines:
        fields = line.split('\t')
        if len(fields) != 3:
            raise ValueError('malformed edge observation')
        formula, kind, raw = fields
        if formula not in expected or formula in seen:
            raise ValueError(f'unknown or duplicate edge: {formula}')
        seen.add(formula)
        want = expected[formula]
        try:
            valid = (kind == 'E' and raw == want) if isinstance(want, str) else (
                kind == 'N' and same_number(float(raw), want))
        except ValueError:
            valid = False
        if not valid:
            raise ValueError(f'wrong typed edge result: {formula}, {kind}, {raw}')
        if reference is not None:
            value = reference(formula)
            if formula in nonfinite:
                if type(value) is not float or not math.isinf(value):
                    raise ValueError(f'changed overflow reference observation: {formula}: {value!r}')
                divergences += 1
            elif isinstance(want, str) or not same_number(float(raw), value):
                raise ValueError(f'edge reference mismatch: {formula}: {raw}, {value!r}')
            else:
                matches += 1
    if seen != set(expected):
        raise ValueError(f'incomplete edge corpus: got {len(seen)}, expected {len(expected)}')
    return matches, divergences


def check_xlsx(path, openpyxl=None):
    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
    cells = {cell.attrib['r']: cell for cell in root.findall('.//m:c', ns)}
    expected = {'A1': (None, '1.005'), 'B1': ('ROUND(A1,2)', '1.01'), 'C1': ('B1*100', '101')}
    if set(cells) != set(expected):
        raise ValueError('unexpected XLSX cells')
    checks = 0
    for ref, (formula, cached) in expected.items():
        fnode = cells[ref].find('m:f', ns)
        vnode = cells[ref].find('m:v', ns)
        if (fnode.text if fnode is not None else None) != formula or vnode is None or vnode.text != cached:
            raise ValueError(f'wrong OOXML formula/cache for {ref}')
        checks += 2
    if openpyxl is not None:
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
        cached = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            for ref, (formula, value) in {'A1': (1.005, 1.005), 'B1': ('=ROUND(A1,2)', 1.01), 'C1': ('=B1*100', 101)}.items():
                if formulas.active[ref].value != formula or cached.active[ref].value != value:
                    raise ValueError(f'openpyxl formula/cache mismatch for {ref}')
                checks += 2
        finally:
            formulas.close()
            cached.close()
    return checks


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--require-reference', action='store_true')
    args = parser.parse_args()
    try:
        import formualizer
        import openpyxl
    except ImportError:
        if args.require_reference:
            raise
        formualizer = openpyxl = None
    def evaluate(formula):
        book = formualizer.Workbook()
        book.add_sheet('Sheet1')
        book.set_formula('Sheet1', 1, 1, formula)
        return book.evaluate_cell('Sheet1', 1, 1)
    reference = evaluate if formualizer is not None else None
    decimal_count, independent_count = compare_rows(
        (ARTIFACTS / 'round-corpus.tsv').read_text(encoding='utf-8').splitlines(), expected_keys(), reference)
    xlsx_count = check_xlsx(ARTIFACTS / 'round.xlsx', openpyxl)
    edges = (ARTIFACTS / 'round-edges.tsv').read_text(encoding='utf-8').splitlines()
    edge_matches, edge_nonfinite = compare_edges(edges, EDGES, reference, NONFINITE_REFERENCE)
    print(f'Python decimal: {decimal_count}/{len(expected_keys())} exact typed binary64 matches')
    print(f'formualizer: {independent_count}/{len(expected_keys())} corpus matches; '
          f'{edge_matches} edge matches, {edge_nonfinite} documented nonfinite divergences' if reference else
          'formualizer unavailable: no independent formula-engine comparison')
    print(f'OOXML/openpyxl checks: {xlsx_count}; openpyxl reads formula caches, does not recalculate')


if __name__ == '__main__':
    main()
