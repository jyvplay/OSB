#!/usr/bin/env python3
"""Actual Python zipfile/ElementTree producers versus separately compiled Kotlin readers. No Android inference."""
import base64, hashlib, io, json, os, random, struct, subprocess, time, zipfile, xml.etree.ElementTree as ET
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'testing/build/r22';rng=random.Random(220022)
requests=[];expected=[];labels=[];exports={};pairs=[]
def b64(b):return base64.b64encode(b).decode()
def case(label,op,raw,result=None,names='a.txt'):
 requests.append(op+'\t'+b64(raw)+'\t'+names);expected.append('REFUSED' if result is None else 'OK\t'+b64(result.encode()));labels.append(label)
def zipped(files,method=8,extra=b'',stream=None):
 b=io.BytesIO() if stream is None else stream
 with zipfile.ZipFile(b,'w',compression=method) as z:
  for name,data in files:
   i=zipfile.ZipInfo(name);i.compress_type=method;i.extra=extra;z.writestr(i,data)
 return b.getvalue()
for i in range(90):
 files=[('p'+str(j)+'.txt',('α😀&<>\r\n'+''.join(rng.choices('abc012',k=rng.randrange(0,5000)))).encode()) for j in range(1+i%7)]
 raw=zipped(files,0 if i%2 else 8)
 with zipfile.ZipFile(io.BytesIO(raw)) as z:
  result='\n'.join(n+':'+b64(z.read(n)) for n in sorted(z.namelist()))
 case('python ZIP '+str(i),'Z',raw,result,','.join(n for n,_ in files))
base=zipped([('a.txt',b'payload')]);cen=base.index(b'PK\x01\x02');end=len(base)-22
for index in [6,8,14,18,22,30,cen+8,cen+10,cen+16,cen+20,cen+24,cen+46,end+4,end+6,end+8,end+10,end+12,end+16,end+20]:
 raw=bytearray(base);raw[index]^=1;case('ZIP bit mutation '+str(index),'Z',raw)
for name,raw in [('trailing',base+b'x'),('truncated',base[:-1]),('unknown',zipped([('other',b'x')])),('zip64-extra',zipped([('a.txt',b'x')],extra=struct.pack('<HHQQ',1,16,1,1))),('expansion',zipped([('a.txt',b'x'*1048577)])),('input-cap',b'x'*2097153)]:case(name,'Z',raw)
class NonSeek(io.BytesIO):
 def seek(self,*args):raise io.UnsupportedOperation()
case('data-descriptor','Z',zipped([('a.txt',b'x')],stream=NonSeek()))
# A valid raw deflate stream followed by extra compressed-member bytes: permissive zipfile accepts it.
raw=bytearray(base);packed=struct.unpack_from('<I',raw,18)[0];start=35
raw[start+packed:start+packed]=b'junk';newcen=cen+4;newend=end+4
struct.pack_into('<I',raw,18,packed+4);struct.pack_into('<I',raw,newcen+20,packed+4);struct.pack_into('<I',raw,newend+16,newcen)
with zipfile.ZipFile(io.BytesIO(raw)) as z:assert z.read('a.txt')==b'payload'
case('deflate trailing data accepted by Python but excluded by policy','Z',raw)
# Actual-output cap, not just declared-size rejection.
raw=bytearray(zipped([('a.txt',b'x'*100000)]));c=raw.index(b'PK\x01\x02');struct.pack_into('<I',raw,22,1);struct.pack_into('<I',raw,c+24,1)
case('lying expansion size','Z',raw)
for i in range(70):
 root=ET.Element('{urn:r22}root',{'a':str(i)});root.text='value α & <> '+str(i)
 for j in range(i%8):ET.SubElement(root,'child').text=str(j)
 raw=ET.tostring(root,encoding='utf-8',xml_declaration=True);parsed=ET.fromstring(raw)
 case('ElementTree '+str(i),'X',raw,'urn:r22|root|'+(parsed.text or '')+'|'+str(len(parsed)))
for label,raw in [('depth32',('<a>'*32+'</a>'*32).encode()),('nodes32768',('<a>'+'<b/>'*32767+'</a>').encode()),('bytecap',('<a>'+'x'*(1048576-7)+'</a>').encode()),('attrs16',('<a '+' '.join('p%d="x"'%i for i in range(16))+'/>').encode())]:
 root=ET.fromstring(raw);case(label,'X',raw,'|a|'+(root.text or '')+'|'+str(len(root)))
for label,raw in [('depth33','<a>'*33+'</a>'*33),('nodes32769','<a>'+'<b/>'*32768+'</a>'),('attrs17','<a '+' '.join('p%d="x"'%i for i in range(17))+'/>'),('doctype',"<!DOCTYPE a [<!ENTITY x SYSTEM 'file:///etc/passwd'>]><a>&x;</a>"),('CDATA','<a><![CDATA[x]]></a>'),('PI','<?other x?><a/>'),('comment','<a><!--x--></a>'),('wrong-encoding',"<?xml version='1.0' encoding='UTF-16'?><a/>"),('duplicate-attribute',"<a x='1' x='2'/>"),('unbound-prefix','<p:a/>'),('wrong-close','<a></b>')]:case(label,'X',raw.encode())

# Real independently maintained openpyxl writes valid XLSX; no Office runtime inference.
from openpyxl import Workbook,load_workbook
from datetime import datetime
def newbook():
 w=Workbook();w.properties.created=datetime(2026,1,1);w.properties.modified=datetime(2026,1,1);return w
def canonical_zip(raw):
 with zipfile.ZipFile(io.BytesIO(raw)) as z:return zipped([(n,z.read(n)) for n in z.namelist()])
for i in range(40):
 wb=newbook();ws=wb.active
 values=[['name','value','private'],['Alpha α',str(i),'OMIT'],['=literal','😀 & < >','line\nnext']]
 for r,row in enumerate(values,1):
  for c,value in enumerate(row,1):ws.cell(r,c,value).data_type='s'
 stream=io.BytesIO();wb.save(stream);raw=canonical_zip(stream.getvalue())
 reference=load_workbook(io.BytesIO(raw),data_only=False).active
 result='\n'.join(','.join(b64(str(v).encode()) for v in row) for row in reference.values)
 case('openpyxl text '+str(i),'W',raw,result)
# Sparse coordinates remain attached to rows/columns, not sequential scanner positions.
wb=newbook();ws=wb.active;ws['C3']='end';ws['A1']=True;ws['B2']=12
stream=io.BytesIO();wb.save(stream);sparse=canonical_zip(stream.getvalue())
case('sparse numeric boolean','W',sparse,'\n'.join(','.join(b64(v.encode()) for v in row) for row in [['1','',''],['','12',''],['','','end']]))
def repack(raw,edit):
 with zipfile.ZipFile(io.BytesIO(raw)) as z:parts={n:z.read(n) for n in z.namelist()}
 edit(parts);return zipped(list(parts.items()))
for label,edit in [
 ('formula',lambda p:p.__setitem__('xl/worksheets/sheet1.xml',p['xl/worksheets/sheet1.xml'].replace(b'<v>12</v>',b'<f>1+1</f><v>12</v>'))),
 ('row-coordinate-disagreement',lambda p:p.__setitem__('xl/worksheets/sheet1.xml',p['xl/worksheets/sheet1.xml'].replace(b'r="B2"',b'r="B3"'))),
 ('hidden-row',lambda p:p.__setitem__('xl/worksheets/sheet1.xml',p['xl/worksheets/sheet1.xml'].replace(b'<row r="2"',b'<row hidden="1" r="2"'))),
 ('unknown-part',lambda p:p.__setitem__('xl/vbaProject.bin',b'x')),
 ('malformed-XML',lambda p:p.__setitem__('xl/worksheets/sheet1.xml',b'<broken>')),
 ('DTD',lambda p:p.__setitem__('xl/worksheets/sheet1.xml',b'<!DOCTYPE worksheet [<!ENTITY x "boom">]>'+p['xl/worksheets/sheet1.xml'])),
 ('external-relationship',lambda p:p.__setitem__('xl/_rels/workbook.xml.rels',p['xl/_rels/workbook.xml.rels'].replace(b'<Relationship ',b'<Relationship TargetMode="External" ',1))),
 ('cell-out-of-column-bound',lambda p:p.__setitem__('xl/worksheets/sheet1.xml',p['xl/worksheets/sheet1.xml'].replace(b'r="C3"',b'r="AG3"'))),
 ('OOXML-escape',lambda p:p.__setitem__('xl/worksheets/sheet1.xml',p['xl/worksheets/sheet1.xml'].replace(b'>end<',b'>_x0041_<'))),
 ('unknown-cell-type',lambda p:p.__setitem__('xl/worksheets/sheet1.xml',p['xl/worksheets/sheet1.xml'].replace(b't="b"',b't="e"'))),
 ('wrong-content-type',lambda p:p.__setitem__('[Content_Types].xml',p['[Content_Types].xml'].replace(b'spreadsheetml.sheet.main+xml',b'spreadsheetml.template.main+xml'))),
 ('relationship-path-traversal',lambda p:p.__setitem__('xl/_rels/workbook.xml.rels',p['xl/_rels/workbook.xml.rels'].replace(b'/xl/worksheets/sheet1.xml',b'/xl/../xl/worksheets/sheet1.xml')))
]:case(label,'W',repack(sparse,edit))
for label,change in [('date-style',lambda w:setattr(w.active['A1'],'number_format','yyyy-mm-dd')),('second-sheet',lambda w:w.create_sheet()),('merge',lambda w:w.active.merge_cells('A1:B1'))]:
 wb=newbook();wb.active['A1']=123;change(wb);stream=io.BytesIO();wb.save(stream);case(label,'W',canonical_zip(stream.getvalue()))

import xlsxwriter
for i in range(30):
 stream=io.BytesIO()
 with xlsxwriter.Workbook(stream,{'in_memory':True}) as wb:
  wb.set_properties({'created':datetime(2026,1,1)});ws=wb.add_worksheet('Shared');ws.write_string(0,0,'=1+1');ws.write_string(2,2,'shared α '+str(i))
 raw=canonical_zip(stream.getvalue())
 reference=load_workbook(io.BytesIO(raw),data_only=False).active
 result='\n'.join(','.join(b64(('' if v is None else str(v)).encode()) for v in row) for row in reference.values)
 case('XlsxWriter shared sparse '+str(i),'W',raw,result)
if '--write-native-fixture' in __import__('sys').argv:
 fixture=ROOT/'app/src/androidTest/assets/r22-shared.xlsx';fixture.parent.mkdir(parents=True,exist_ok=True);fixture.write_bytes(raw)

# Source worksheet names, package author metadata and omitted cells must not affect new export bytes.
for pair in range(15):
 indices=[]
 for variant in ['Alice','Bob']:
  wb=newbook();wb.active.title=variant;wb.properties.creator=variant
  wb.active.append([variant,'=1+1']);wb.active['B1'].data_type='s'
  wb.active.append(['PRIVATE-'+variant,'keep '+str(pair)])
  stream=io.BytesIO();wb.save(stream);raw=canonical_zip(stream.getvalue())
  indices.append(len(requests));exports[len(requests)]=[['=1+1'],['keep '+str(pair)]]
  requests.append('E\t'+b64(raw)+'\t1');expected.append('EXPORT');labels.append('metadata/omitted-cell export '+str(pair)+' '+variant)
 pairs.append(indices)
for count,unit,admit in [(4096,'x',True),(4097,'x',False),(2048,'😀',True),(2049,'😀',False)]:
 wb=newbook();wb.active['A1']=unit*count;stream=io.BytesIO();wb.save(stream)
 case('UTF16 cell boundary '+str(count)+unit,'W',canonical_zip(stream.getvalue()),b64((unit*count).encode()) if admit else None)

for value in [0.5,-12.25,1e-7,1e20]:
 stream=io.BytesIO()
 with xlsxwriter.Workbook(stream,{'in_memory':True}) as wb:
  wb.set_properties({'created':datetime(2026,1,1)});wb.add_worksheet().write_number(0,0,value)
 raw=canonical_zip(stream.getvalue())
 with zipfile.ZipFile(io.BytesIO(raw)) as z:lexical=ET.fromstring(z.read('xl/worksheets/sheet1.xml')).find('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v').text
 case('numeric lexical '+str(value),'W',raw,b64(lexical.encode()))
cp=os.pathsep.join(str(p) for p in sorted(OUT.glob('*.jar')));java=str(Path(os.environ['JAVA_HOME'])/'bin/java') if 'JAVA_HOME' in os.environ else 'java'
payload='\n'.join(requests)+'\n';start=time.monotonic();r=subprocess.run([java,'-cp',cp,'dev.fusion.testing.R22Oracle'],input=payload,text=True,capture_output=True,check=True,timeout=120)
actual=r.stdout.splitlines();assert len(actual)==len(expected),(len(actual),len(expected),r.stderr)
books={}
for i,(label,a,e) in enumerate(zip(labels,actual,expected)):
 if e=='EXPORT':
  assert a.startswith('OK\t'),(label,a)
  book=base64.b64decode(base64.b64decode(a.split('\t')[1]));books[i]=book
  sheet=load_workbook(io.BytesIO(book),data_only=False).active
  assert [list(row) for row in sheet.values]==exports[i],label
  assert all(c.data_type=='s' for row in sheet for c in row),label
 else:assert a==e,(label,a[:150],e[:150])
for a,b in pairs:assert books[a]==books[b],(labels[a],labels[b])
report={'python':__import__('sys').version.split()[0],'openpyxl':__import__('openpyxl').__version__,'xlsxwriter':xlsxwriter.__version__,'script_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'metadata_note':'Producer modified timestamps can vary; seed fixes cell corpus, not every input byte.','seed':220022,'requests':len(requests),'accepted':sum(x.startswith('OK') or x=='EXPORT' for x in expected),'refused':expected.count('REFUSED'),'equal_byte_pairs':len(pairs),'request_sha256':hashlib.sha256(payload.encode()).hexdigest(),'output_sha256':hashlib.sha256(r.stdout.encode()).hexdigest(),'seconds_observed':round(time.monotonic()-start,3),'labels':labels,'scope':'Authored corpora; actual Python ZIP/XML producers versus actual Kotlin implementations on JDK, not Android/Office runtime or blind audit.'}
(ROOT/'docs/verification/r22-interop.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps({k:v for k,v in report.items() if k!='labels'}))
