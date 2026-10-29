#!/usr/bin/env python3
"""Independent Python + openpyxl + zipfile + hashlib differential reference oracle for R45.

Executes compiled Kotlin R45 production kernels on the host JVM and verifies byte/numeric
agreement against independent Python implementations and held-out third-party parsers
(`openpyxl` in both `data_only=False` and `data_only=True` modes, `zipfile` + `ElementTree`,
`hmac`/`hashlib`, and analytical math).
"""
from __future__ import annotations

import argparse
import hashlib
import hmac
import io
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

_HOST_PY = Path(os.environ.get("FUSION_HOST_CACHE", str(Path.home() / ".cache/fusion-host-r23"))) / "python"
if _HOST_PY.is_dir() and str(_HOST_PY) not in sys.path:
    sys.path.insert(0, str(_HOST_PY))

import openpyxl  # installed by bootstrap_host_tools.py in ~/.cache/fusion-host-r23/python

import run_r45_contracts

ROOT = Path(__file__).resolve().parents[1]
BRIDGE_DIR = ROOT / "testing/build/r45-ref-bridge"

KOTLIN_PROBE_SRC = r"""
package dev.fusion.testing

import dev.fusion.calc.FinMath
import dev.fusion.designsystem.SpringStep
import dev.fusion.files.MerkleTree
import dev.fusion.index.Bm25Rank
import dev.fusion.maps.UtmGrid
import dev.fusion.native.RleBitset
import dev.fusion.ooxml.DocxPackage
import dev.fusion.ooxml.FormulaEval
import dev.fusion.ooxml.FormulaXlsx
import dev.fusion.vault.HkdfExpand
import java.nio.charset.StandardCharsets
import java.nio.file.Files
import java.nio.file.Paths

object R45RefProbe {
    private fun hex(b: ByteArray): String = b.joinToString("") { "%02x".format(it.toInt() and 0xFF) }

    @JvmStatic
    fun main(args: Array<String>) {
        val workDir = Paths.get(args[0])
        val openpyxlIn = Paths.get(args[1])

        // 1. Emit Kotlin FormulaXlsx .xlsx for openpyxl dual-mode inspection
        val xlsxBytes = FormulaXlsx.write(
            "Q4Model",
            linkedMapOf(
                "A1" to FormulaEval.CellInput.TextLiteral("Revenue"),
                "B1" to FormulaEval.CellInput.NumberLiteral(1250.25),
                "B2" to FormulaEval.CellInput.NumberLiteral(749.75),
                "B3" to FormulaEval.CellInput.Formula("=SUM(B1:B2)"),
                "B4" to FormulaEval.CellInput.Formula("=IF(B3>=2000, \"PASS\", \"FAIL\")"),
                "B5" to FormulaEval.CellInput.Formula("=ROUND(B3*15%, 2)"),
                "B6" to FormulaEval.CellInput.Formula("=1/0"),
                "B7" to FormulaEval.CellInput.Formula("=AND(B3=2000, B1>1000)"),
            ),
        )
        Files.write(workDir.resolve("kotlin_out.xlsx"), xlsxBytes)

        // 2. Read openpyxl-authored .xlsx (with sharedStrings + formulas) and recalculate in Kotlin
        val foreignLoaded = FormulaXlsx.read(Files.readAllBytes(openpyxlIn))
        val foreignEval = foreignLoaded.recalculate()

        // 3. Emit Kotlin DocxPackage .docx for zipfile + ElementTree inspection
        val docxBytes = DocxPackage.write(
            DocxPackage.Document(
                title = "Audit",
                blocks = listOf(
                    DocxPackage.Block.Heading(1, listOf(DocxPackage.Run("Section One"))),
                    DocxPackage.Block.Paragraph(
                        listOf(DocxPackage.Run("Plain "), DocxPackage.Run("bold", bold = true)),
                    ),
                    DocxPackage.Block.ListItem(ordered = true, level = 0, runs = listOf(DocxPackage.Run("First"))),
                    DocxPackage.Block.Table(listOf(listOf("K", "V"), listOf("alpha", "42"))),
                ),
            ),
        )
        Files.write(workDir.resolve("kotlin_out.docx"), docxBytes)

        // 4. HKDF vectors
        val hkdfHex = (1..6).map { n ->
            val ikm = ByteArray(16 + n) { (it * 7 + n).toByte() }
            val salt = if (n % 2 == 0) ByteArray(0) else ByteArray(12) { (it + 3).toByte() }
            val info = "r45-ctx-$n".toByteArray(StandardCharsets.UTF_8)
            hex(HkdfExpand.deriveKey(ikm, salt, info, 20 + n * 9))
        }

        // 5. Merkle roots for sizes 0..10
        val merkleHex = (0..10).map { sz ->
            val leaves = (0 until sz).map { idx -> "leaf-$sz-$idx".toByteArray(StandardCharsets.UTF_8) }
            MerkleTree.toHex(MerkleTree.rootHash(leaves))
        }

        // 6. FinMath + SpringStep + UtmGrid + RleBitset
        val pmtVal = FinMath.pmt(0.005, 240.0, 150000.0, 0.0, 0)
        val npvVal = FinMath.npv(0.08, listOf(500.0, 600.0, 700.0))
        val spUnder = SpringStep.evaluate(SpringStep.Spec(256.0, 0.6), 0.0, 0.0, 10.0, 0.15)
        val spOver = SpringStep.evaluate(SpringStep.Spec(256.0, 1.5), 0.0, 0.0, 10.0, 0.15)
        val utmOslo = UtmGrid.fromLatLon(59.9139, 10.7522)
        val rleHex = hex(RleBitset.encode(RleBitset.ofInts(listOf(0, 1, 2, 10, 11, 500))))

        val lines = buildList {
            add("foreign_sheet=${foreignLoaded.sheetName}")
            add("foreign_C1=${foreignEval.values["C1"]?.displayText()}")
            add("foreign_C2=${foreignEval.values["C2"]?.displayText()}")
            add("hkdf=${hkdfHex.joinToString(",")}")
            add("merkle=${merkleHex.joinToString(",")}")
            add("pmt=$pmtVal")
            add("npv=$npvVal")
            add("spring_under=${spUnder.value},${spUnder.velocity}")
            add("spring_over=${spOver.value},${spOver.velocity}")
            add("utm_oslo=${utmOslo.zone},${utmOslo.band},${utmOslo.eastingMeters},${utmOslo.northingMeters}")
            add("rle_hex=$rleHex")
        }
        Files.write(workDir.resolve("probe_out.txt"), lines, StandardCharsets.UTF_8)
    }
}
"""


def py_hkdf_sha256(ikm: bytes, salt: bytes, info: bytes, length: int) -> bytes:
    eff_salt = salt if salt else b"\x00" * 32
    prk = hmac.new(eff_salt, ikm, hashlib.sha256).digest()
    t = b""
    okm = bytearray()
    counter = 1
    while len(okm) < length:
        t = hmac.new(prk, t + info + bytes([counter]), hashlib.sha256).digest()
        okm.extend(t)
        counter += 1
    return bytes(okm[:length])


def py_merkle_root(leaves: list[bytes]) -> str:
    if not leaves:
        return hashlib.sha256(b"").hexdigest()

    def mth(sub: list[bytes]) -> bytes:
        n = len(sub)
        if n == 1:
            return hashlib.sha256(b"\x00" + sub[0]).digest()
        k = 1 << ((n - 1).bit_length() - 1)
        return hashlib.sha256(b"\x01" + mth(sub[:k]) + mth(sub[k:])).digest()

    return mth(leaves).hex()


def py_spring(stiffness: float, zeta: float, target: float, dt: float) -> tuple[float, float]:
    x0 = -target
    v0 = 0.0
    w0 = math.sqrt(stiffness)
    if zeta < 1.0:
        wd = w0 * math.sqrt(1.0 - zeta * zeta)
        env = math.exp(-zeta * w0 * dt)
        c1, c2 = x0, (v0 + zeta * w0 * x0) / wd
        x = env * (c1 * math.cos(wd * dt) + c2 * math.sin(wd * dt))
        v = -zeta * w0 * x + env * (-c1 * wd * math.sin(wd * dt) + c2 * wd * math.cos(wd * dt))
        return target + x, v
    gamma = w0 * math.sqrt(zeta * zeta - 1.0)
    r1, r2 = -zeta * w0 + gamma, -zeta * w0 - gamma
    c2 = (v0 - r1 * x0) / (r2 - r1)
    c1 = x0 - c2
    x = c1 * math.exp(r1 * dt) + c2 * math.exp(r2 * dt)
    v = c1 * r1 * math.exp(r1 * dt) + c2 * r2 * math.exp(r2 * dt)
    return target + x, v


def verify() -> dict[str, object]:
    kotlinc, java_bin, stdlib, env = run_r45_contracts.resolve_toolchain()
    jars, _ = run_r45_contracts.compile_all(workers=2)
    BRIDGE_DIR.mkdir(parents=True, exist_ok=True)
    probe_kt = BRIDGE_DIR / "R45RefProbe.kt"
    probe_kt.write_text(KOTLIN_PROBE_SRC.strip() + "\n")
    probe_jar = BRIDGE_DIR / "r45-ref-probe.jar"

    cp = os.pathsep.join(str(j) for j in jars)
    subprocess.run(
        [str(kotlinc), "-Werror", "-cp", cp, str(probe_kt), "-d", str(probe_jar)],
        cwd=ROOT, env=env, check=True,
    )

    with tempfile.TemporaryDirectory(prefix="r45-diff-") as tmp:
        tmp_path = Path(tmp)
        foreign_xlsx = tmp_path / "openpyxl_in.xlsx"

        # Author an .xlsx workbook via real Python openpyxl (uses sharedStrings + formula cells)
        wb_in = openpyxl.Workbook()
        ws_in = wb_in.active
        ws_in.title = "OpenPyxlSrc"
        ws_in["A1"] = "Label"
        ws_in["B1"] = 40
        ws_in["B2"] = 60
        ws_in["C1"] = "=SUM(B1:B2)*2"
        ws_in["C2"] = '=CONCAT(A1, "-", C1)'
        wb_in.save(foreign_xlsx)

        run_cp = os.pathsep.join([*(str(j) for j in jars), str(probe_jar), str(stdlib)])
        subprocess.run(
            [str(java_bin), "-cp", run_cp, "dev.fusion.testing.R45RefProbe", str(tmp_path), str(foreign_xlsx)],
            cwd=ROOT, env=env, check=True,
        )

        # Verify Kotlin FormulaXlsx output in openpyxl data_only=False AND data_only=True!
        wb_f = openpyxl.load_workbook(tmp_path / "kotlin_out.xlsx", data_only=False)
        assert wb_f.sheetnames == ["Q4Model"], wb_f.sheetnames
        ws_f = wb_f["Q4Model"]
        assert ws_f["A1"].value == "Revenue"
        assert ws_f["B3"].value == "=SUM(B1:B2)", ws_f["B3"].value
        assert ws_f["B5"].value == "=ROUND(B3*15%, 2)", ws_f["B5"].value

        wb_v = openpyxl.load_workbook(tmp_path / "kotlin_out.xlsx", data_only=True)
        ws_v = wb_v["Q4Model"]
        assert ws_v["B3"].value == 2000, ws_v["B3"].value
        assert ws_v["B4"].value == "PASS", ws_v["B4"].value
        assert abs(float(ws_v["B5"].value) - 300.0) < 1e-9
        assert ws_v["B6"].value == "#DIV/0!", ws_v["B6"].value
        assert ws_v["B7"].value is True, ws_v["B7"].value

        # Verify Kotlin DocxPackage output via Python zipfile + ElementTree
        with zipfile.ZipFile(tmp_path / "kotlin_out.docx", "r") as zf:
            names = set(zf.namelist())
            required = {
                "[Content_Types].xml", "_rels/.rels", "word/document.xml",
                "word/styles.xml", "word/numbering.xml", "word/_rels/document.xml.rels",
            }
            assert required.issubset(names), names
            for info in zf.infolist():
                assert info.date_time == (1980, 1, 1, 0, 0, 0), f"non-zero-epoch {info.date_time}"
            doc_root = ET.fromstring(zf.read("word/document.xml"))
            w_ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            texts = [t.text or "" for t in doc_root.findall(".//w:t", w_ns)]
            assert "Section One" in texts and "bold" in texts and "42" in texts, texts

        # Parse Kotlin scalar/digest probe outputs and compare against Python reference math
        kv: dict[str, str] = {}
        for line in (tmp_path / "probe_out.txt").read_text().splitlines():
            k, v = line.split("=", 1)
            kv[k] = v

        assert kv["foreign_sheet"] == "OpenPyxlSrc"
        assert kv["foreign_C1"] == "200", kv["foreign_C1"]
        assert kv["foreign_C2"] == "Label-200", kv["foreign_C2"]

        kt_hkdfs = kv["hkdf"].split(",")
        for idx, got_hex in enumerate(kt_hkdfs, start=1):
            ikm = bytes(((i * 7 + idx) & 0xFF) for i in range(16 + idx))
            salt = b"" if idx % 2 == 0 else bytes(((i + 3) & 0xFF) for i in range(12))
            info = f"r45-ctx-{idx}".encode("utf-8")
            exp_hex = py_hkdf_sha256(ikm, salt, info, 20 + idx * 9).hex()
            assert got_hex == exp_hex, f"HKDF mismatch at {idx}: {got_hex} != {exp_hex}"

        kt_merkles = kv["merkle"].split(",")
        for sz, got_m in enumerate(kt_merkles):
            leaves = [f"leaf-{sz}-{i}".encode("utf-8") for i in range(sz)]
            exp_m = py_merkle_root(leaves)
            assert got_m == exp_m, f"Merkle mismatch at size={sz}: {got_m} != {exp_m}"

        r, nper, pv = 0.005, 240.0, 150000.0
        exp_pmt = -(pv * ((1.0 + r) ** nper) * r) / (((1.0 + r) ** nper) - 1.0)
        assert abs(float(kv["pmt"]) - exp_pmt) < 1e-8
        exp_npv = 500.0 / 1.08 + 600.0 / (1.08 ** 2) + 700.0 / (1.08 ** 3)
        assert abs(float(kv["npv"]) - exp_npv) < 1e-9

        ux, uv = [float(x) for x in kv["spring_under"].split(",")]
        eux, euv = py_spring(256.0, 0.6, 10.0, 0.15)
        assert abs(ux - eux) < 1e-9 and abs(uv - euv) < 1e-8

        ox, ov = [float(x) for x in kv["spring_over"].split(",")]
        eox, eov = py_spring(256.0, 1.5, 10.0, 0.15)
        assert abs(ox - eox) < 1e-9 and abs(ov - eov) < 1e-8

        zone_s, band_s, east_s, north_s = kv["utm_oslo"].split(",")
        assert int(zone_s) == 32 and band_s == "V"
        assert 597_000.0 < float(east_s) < 599_000.0 and 6_642_000.0 < float(north_s) < 6_644_000.0

        # RleBitset wire check: 'R'(0x52) 'B'(0x42) ver=1 count=3:
        # run0(start=0,len=3)->(0,2); run1(start=10,len=2,prevEnd=3)->(6,1); run2(start=500,len=1,prevEnd=12)->(487=0xe7,0x03, 0)
        assert kv["rle_hex"] == "5242010300020601e70300", kv["rle_hex"]

    return {
        "openpyxl_dual_mode_xlsx": "PASS",
        "openpyxl_to_kotlin_xlsx_ingest": "PASS",
        "docx_opc_zero_epoch_xml": "PASS",
        "hkdf_rfc5869_vectors": 6,
        "merkle_rfc6962_trees": 11,
        "analytical_physics_geodesy_tvm": "PASS",
    }


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--verify", action="store_true", help="Run differential verification")
    ap.add_argument("--json", action="store_true", help="Print JSON summary")
    args = ap.parse_args(argv)
    summary = verify()
    if args.json:
        print(json.dumps(summary, indent=2))
    else:
        print("R45 reference differential oracle PASS:", json.dumps(summary))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
