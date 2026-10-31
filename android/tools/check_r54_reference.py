#!/usr/bin/env python3
"""Independent ZIP/XML checks plus optional openpyxl/python-docx readers for R54 artifacts.

Run after tools/run_r54_contracts.py. Readers verify structure and cached values; neither reader
calculates formulas or proves the behavior of Word, Excel, Android, or a document provider.
"""
from pathlib import Path
import argparse
import sys
import zipfile
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / 'testing/build/r54/artifacts'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--require-office-libraries', action='store_true')
    args = parser.parse_args()
    checks = 0

    def check(condition, reason):
        nonlocal checks
        checks += 1
        if not condition:
            raise AssertionError(reason)

    expected_paragraphs = ['  First\tcolumn & <quoted> 🙂', 'Second', '', 'last', '']
    with zipfile.ZipFile(ART / 'writer.docx') as package:
        check(package.testzip() is None, 'DOCX ZIP CRC')
        names = set(package.namelist())
        check('word/document.xml' in names, 'DOCX main part')
        check(not any('vba' in n.lower() or n.startswith('docProps/') for n in names), 'DOCX has no macro or document properties')
        document = ET.fromstring(package.read('word/document.xml'))
        paragraphs = []
        for p in document.find(W + 'body').findall(W + 'p'):
            chunks = []
            for node in p.iter():
                if node.tag == W + 't':
                    chunks.append(node.text or '')
                elif node.tag == W + 'tab':
                    chunks.append('\t')
                elif node.tag == W + 'br':
                    chunks.append('\n')
            paragraphs.append(''.join(chunks))
        check(paragraphs == expected_paragraphs, 'DOCX paragraph and tab content')
        check(not any(b'TargetMode="External"' in package.read(name) for name in names if name.endswith('.rels')), 'DOCX has no external relationships')

    with zipfile.ZipFile(ART / 'sheet.xlsx') as package:
        check(package.testzip() is None, 'XLSX ZIP CRC')
        names = set(package.namelist())
        check('xl/worksheets/sheet1.xml' in names, 'XLSX worksheet part')
        check(not any('vba' in n.lower() or 'externalLink' in n for n in names), 'XLSX has no macro or external-link parts')
        sheet = ET.fromstring(package.read('xl/worksheets/sheet1.xml'))
        cells = {c.attrib['r']: c for c in sheet.iter(S + 'c')}
        check(cells['B1'].find(S + 'f').text == 'SUM(A1:A2)', 'formula is serialized')
        check(float(cells['B1'].find(S + 'v').text) == 2 + 3, 'cached arithmetic result')
        check(cells['B2'].find(S + 'v').text == 'yes', 'cached text result')
        check(cells['C1'].attrib.get('t') == 'inlineStr' and cells['C1'].find('.//' + S + 't').text == '=1+1', 'formula-looking input is literal')
        check(cells['C2'].find('.//' + S + 't').text == '\t spaced & <data>\r\n🙂', 'XML preserves CRLF and Unicode')
        check(cells['D1'].attrib.get('t') == 'b' and cells['D1'].find(S + 'v').text == '1', 'boolean type')
        check(cells['E1'].attrib.get('t') == 'e' and cells['E1'].find(S + 'v').text == '#DIV/0!', 'formula error cache')

    used = []
    try:
        import openpyxl
    except ImportError:
        if args.require_office_libraries:
            raise
    else:
        formulas = openpyxl.load_workbook(ART / 'sheet.xlsx', data_only=False).active
        cached = openpyxl.load_workbook(ART / 'sheet.xlsx', data_only=True).active
        check(formulas['B1'].value == '=SUM(A1:A2)' and formulas['B1'].data_type == 'f', 'openpyxl formula')
        check(cached['B1'].value == 5 and cached['B2'].value == 'yes', 'openpyxl cached results')
        check(formulas['C1'].value == '=1+1' and formulas['C1'].data_type == 's', 'openpyxl literal formula-like text')
        check(formulas['C2'].value == '\t spaced & <data>\r\n🙂', 'openpyxl text fidelity')
        check(cached['D1'].value is True and cached['E1'].value == '#DIV/0!', 'openpyxl typed boolean/error')
        used.append('openpyxl ' + openpyxl.__version__)
    try:
        import docx
    except ImportError:
        if args.require_office_libraries:
            raise
    else:
        document = docx.Document(ART / 'writer.docx')
        check([p.text for p in document.paragraphs] == expected_paragraphs, 'python-docx paragraphs and tab')
        used.append('python-docx ' + docx.__version__)
    print(f'R54 reference: {checks} checks passed; independent readers: {", ".join(used) if used else "stdlib ZIP/XML only"}.')
    print('Formula caches were checked against explicit expected results; readers did not calculate them.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
