#!/usr/bin/env python3
"""Inspect R55 XLSX and compare a finite formula corpus with independent host tools.

Requires R55 artifacts. With --require-reference, openpyxl 3.1.5 and formualizer 0.13.0
are mandatory. Formualizer is a separate implementation, NOT Excel or an unquestionable
oracle. Known differences and bounded refusals are counted separately, never as matches.
No reference package is bundled with the Android application.
"""
from pathlib import Path
import argparse
import base64
import importlib.metadata
import math
import sys
import zipfile
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / 'testing/build/r55/artifacts'
S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
ERRORS = {'Div': '#DIV/0!', 'Value': '#VALUE!', 'Num': '#NUM!', 'Name': '#NAME?',
          'Ref': '#REF!', 'NA': '#N/A', 'Null': '#NULL!'}
# The reference ignores direct nonnumeric text for these three aggregates. Do not
# automatically "fix" production to match it: Excel's documented behavior is #VALUE!.
KNOWN_DIFFERENCES = {'=AVERAGE("x",2)', '=MIN("x",2)', '=MAX("x",2)'}
EXPECTED_REFUSALS = {'=AVERAGE(1e308,1e308)'}


def same(kind, actual, reference):
    """Typed comparison; booleans must not compare equal to Python's integer 0/1."""
    if kind == 'N':
        try:
            return type(reference) in (int, float) and math.isclose(float(actual), reference, rel_tol=1e-12, abs_tol=0.0)
        except (ValueError, OverflowError):
            return False
    if kind == 'T':
        return type(reference) is str and actual == reference
    if kind == 'B':
        return type(reference) is bool and actual == str(reference).upper()
    if kind == 'Z':
        return reference is None
    if kind == 'E':
        return type(reference) is dict and reference.get('type') == 'Error' and ERRORS.get(reference.get('kind')) == actual
    return False


def reference_value(module, formula):
    workbook = module.Workbook()
    workbook.add_sheet('Sheet1')
    for row, value in ((1, 3), (2, -2), (3, '5'), (4, True), (8, 'Hello'), (9, 'HELLO'), (10, ' ')):
        workbook.set_value('Sheet1', row, 1, value)
    workbook.set_formula('Sheet1', 6, 1, '=1/0')
    workbook.set_formula('Sheet1', 7, 1, '=""')
    workbook.set_formula('Sheet1', 1, 2, formula)
    return workbook.evaluate_cell('Sheet1', 1, 2)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--require-reference', action='store_true')
    args = parser.parse_args()
    checks = 0

    def check(condition, message):
        nonlocal checks
        checks += 1
        if not condition:
            raise AssertionError(message)

    with zipfile.ZipFile(ART / 'blank-references.xlsx') as package:
        check(package.testzip() is None, 'XLSX ZIP CRC')
        sheet = ET.fromstring(package.read('xl/worksheets/sheet1.xml'))
        cells = {node.attrib['r']: node for node in sheet.iter(S + 'c')}
        check('A1' not in cells, 'source is truly unoccupied')
        check(cells['B1'].find(S + 'f').text == 'A1', 'blank reference formula stored')
        for ref in ('B1', 'B2', 'B3'):
            check(cells[ref].attrib.get('t', 'n') == 'n' and float(cells[ref].find(S + 'v').text) == 0.0,
                  f'{ref} caches a numeric zero')
        check(cells['B4'].attrib.get('t') == 'str' and cells['B4'].find(S + 'v') is not None,
              'empty-string result is tagged as string, not a blank source')
        check(float(cells['C4'].find(S + 'v').text) == 3 and float(cells['C5'].find(S + 'v').text) == 4,
              'COUNT and COUNTA caches')

    used = []
    try:
        import openpyxl
    except ImportError:
        if args.require_reference:
            raise
    else:
        if args.require_reference:
            check(openpyxl.__version__ == '3.1.5', 'unexpected openpyxl version')
        formulas = openpyxl.load_workbook(ART / 'blank-references.xlsx', data_only=False).active
        cached = openpyxl.load_workbook(ART / 'blank-references.xlsx', data_only=True).active
        check(formulas['B1'].value == '=A1' and formulas['B4'].value == '=IFERROR(1/0,"")', 'openpyxl formula expressions')
        check(cached['A1'].value is None and all(cached[r].value == 0 for r in ('B1', 'B2', 'B3')), 'openpyxl numeric caches')
        check(cached['C4'].value == 3 and cached['C5'].value == 4, 'openpyxl downstream count caches')
        check(cached['C6'].value is False and cached['C7'].value is True, 'openpyxl comparison caches')
        # openpyxl reports an empty cached string as None. The XML type check above,
        # not this reader, is the evidence that an explicit string was serialized.
        check(cached['B4'].value is None, 'record openpyxl empty-cache interpretation')
        multi_formulas = openpyxl.load_workbook(ART / 'blank-multisheet.xlsx', data_only=False)
        multi_cached = openpyxl.load_workbook(ART / 'blank-multisheet.xlsx', data_only=True)
        check(multi_formulas.sheetnames == ['Source', 'Summary'] and
              multi_formulas['Summary']['B3'].value == '=COUNT(Source!B1)', 'multisheet formula expression')
        values = [multi_cached[sheet][cell].value for sheet, cell in
                  [('Source', 'B1'), ('Summary', 'B1'), ('Summary', 'B2'), ('Summary', 'B3'),
                   ('Summary', 'B4'), ('Summary', 'B5')]]
        check(multi_cached['Source']['A1'].value is None and
              all(type(value) in (int, float) and value == expected for value, expected in
                  zip(values, [0, 0, 0, 1, 0, 1])), 'multisheet typed blank and count caches')
        used.append('openpyxl ' + openpyxl.__version__)

    try:
        import formualizer
    except ImportError:
        if args.require_reference:
            raise
        print('Formualizer unavailable: formula differential was NOT run.')
    else:
        version = importlib.metadata.version('formualizer')
        if args.require_reference:
            check(version == '0.13.0', 'unexpected formualizer version')
        rows = []
        for line in (ART / 'formula-corpus.tsv').read_text(encoding='utf-8').splitlines():
            fields = line.split('\t')
            check(len(fields) == 3, 'malformed corpus row')
            formula, actual = (base64.b64decode(fields[i], validate=True).decode('utf-8') for i in (0, 2))
            rows.append((formula, fields[1], actual))
        check(len(rows) >= 900 and len({row[0] for row in rows}) == len(rows), 'corpus size and uniqueness')
        matches = 0
        differences = set()
        refusals = set()
        unexpected = []
        for formula, kind, actual in rows:
            ref = reference_value(formualizer, formula)
            if same(kind, actual, ref):
                matches += 1
            elif formula in KNOWN_DIFFERENCES and kind == 'E' and actual == '#VALUE!' and ref == 2.0:
                differences.add(formula)
            elif formula in EXPECTED_REFUSALS and kind == 'R' and 'non-finite' in actual and ref == {'type': 'Error', 'kind': 'Num'}:
                refusals.add(formula)
            else:
                unexpected.append((formula, kind, actual, ref))
        check(differences == KNOWN_DIFFERENCES, f'known differences changed: {differences}')
        check(refusals == EXPECTED_REFUSALS, f'known refusals changed: {refusals}')
        check(not unexpected, f'unexpected formula differences: {unexpected[:12]}')
        print(f'Independent formula comparison: {len(rows)} cases; {matches} typed matches, '
              f'{len(differences)} known reference disagreements, {len(refusals)} bounded refusal(s), '
              f'{len(unexpected)} unexpected disagreements.')
        used.append('formualizer ' + version)
    print(f'R55 XLSX/corpus checks: {checks}; independent tools: {", ".join(used) if used else "stdlib ZIP/XML only"}.')
    print('openpyxl reads caches, not calculations; formualizer is independent but is neither Excel nor LibreOffice.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
