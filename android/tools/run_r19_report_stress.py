#!/usr/bin/env python3
"""Postimplementation adversarial corpus against the frozen actual producer, not a pre-registered blind holdout."""
import base64,datetime,hashlib,io,json,os,random,resource,struct,subprocess,uuid,zlib
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'testing/build/r19';rng=random.Random(190919)
for line in (OUT/'result.txt').read_text().splitlines():
    if len(line)>66 and line[64:66]=='  ':assert hashlib.sha256((ROOT/line[66:]).read_bytes()).hexdigest()==line[:64]
requests=[];cases=[];input_sizes=[]
for i in range(30):
    tasks=[]
    for j in range(0 if i==0 else 128):
        if i==1:title="'"*256
        elif i==2:title='界'*256
        elif i==3:title='🙂'*128
        else:
            title='=literal '
            while len(title.encode('utf-16-le'))//2<254:title+=rng.choice(['&','<','>','"',"'",'界','🙂',' '])
        date=[None,datetime.date(1,1,1),datetime.date(9999,12,31),datetime.date(2028,2,29)][j%4]
        tasks.append((str(uuid.UUID(int=j+1)),title,date,j%2==0,j%3==0))
    raw=struct.pack('>4sqI',b'FTS2',19,len(tasks))
    for identity,title,date,done,archived in tasks:
        text=title.encode();raw+=identity.encode()+struct.pack('>BH',int(done)+2*int(archived),len(text))+text+struct.pack('>q',(date-datetime.date(1970,1,1)).days if date else -(2**63))
    wire=raw+struct.pack('>I',zlib.crc32(raw));input_sizes.append(len(wire));requests.append('T\t'+base64.b64encode(wire).decode());cases.append(tasks)
java=str(Path(os.environ['JAVA_HOME'])/'bin/java');cp=os.pathsep.join(str(OUT/f) for f in ['checks.jar','TaskChanges.jar','LiteralWorkbook.jar'])
r=subprocess.run([java,'-cp',cp,'dev.fusion.testing.R19Oracle'],input='\n'.join(requests)+'\n',capture_output=True,text=True,check=True)
outputs=r.stdout.splitlines();assert len(outputs)==len(cases);sizes=[]
for encoded,tasks in zip(outputs,cases):
    raw=base64.b64decode(encoded);assert len(raw)<=512*1024;sizes.append(len(raw));book=openpyxl.load_workbook(io.BytesIO(raw))
    assert book.sheetnames==['Active','Archived']
    for archived,name in [(False,'Active'),(True,'Archived')]:
        rows=list(book[name].iter_rows(min_row=2));wanted=[t for t in tasks if t[4]==archived];assert len(rows)==len(wanted),(len(rows),len(wanted))
        for row,t in zip(rows,wanted):
            assert row[0].value==t[0] and row[1].value==t[1] and row[1].data_type=='s'
            assert row[2].value==str(t[2]) if t[2] else row[2].value in (None,'')
            assert row[3].value==t[3] and row[4].value==t[4]
artifact={'result':'PASS','cases':30,'seed':190919,'input_bytes_max':max(input_sizes),'output_bytes_min':min(sizes),'output_bytes_max':max(sizes),
 'producer_child_peak_rss_kib_linux':resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss,
 'scope':'Actual host Java process and openpyxl, thirty postimplementation report corpora. RSS is one host observation, not a bound/Android physical budget, and excludes Python reader memory. Not an independently authored blind holdout.',
 'requests_sha256':hashlib.sha256(('\n'.join(requests)+'\n').encode()).hexdigest(),'outputs_sha256':hashlib.sha256(r.stdout.encode()).hexdigest(),
 'sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [Path(__file__),OUT/'checks.jar',OUT/'TaskChanges.jar',OUT/'LiteralWorkbook.jar']}}
(ROOT/'docs/verification/r19-report-stress.json').write_text(json.dumps(artifact,indent=2)+'\n');print(json.dumps(artifact))
