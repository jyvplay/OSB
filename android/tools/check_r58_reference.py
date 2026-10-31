#!/usr/bin/env python3
"""Independent CSV/ZIP/XML/openpyxl inspection of actual R58 artifacts; not Excel execution."""
import csv
import io
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / 'testing/build/r58/artifacts'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
EXPECTED = [['  hello  ', '=1+1', 'line 1\r\nline 2'], ['Q"A', '', '🙂\tT']]


def validate(directory: Path = ARTIFACTS) -> int:
    # csv.reader is a separately implemented parser, not the production parser.
    from openpyxl import load_workbook
    checks = 0

    def require(ok, message):
        nonlocal checks
        checks += 1
        if not ok:
            raise AssertionError(message)

    with (directory / 'input.tsv').open('r', encoding='utf-8', newline='') as stream:
        rows = list(csv.reader(stream, delimiter='\t', quotechar='"', strict=True))
    require(rows == EXPECTED, f'unexpected independent TSV parse: {rows!r}')
    data = (directory / 'pasted.xlsx').read_bytes()
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        names = package.namelist()
        require(len(names) == len(set(names)), 'duplicate ZIP members')
        require('xl/worksheets/sheet1.xml' in names, 'missing sheet')
        require(not any('externallink' in n.lower() or 'vbaproject' in n.lower() for n in names), 'unexpected active external part')
        root = ET.fromstring(package.read('xl/worksheets/sheet1.xml'))
        cells = {c.attrib['r']: c for c in root.findall('.//s:c', NS)}
        require(len(cells) == len(root.findall('.//s:c', NS)), 'duplicate cell addresses')
        expected_cells = {'A1': EXPECTED[0][0], 'B1': EXPECTED[0][1], 'C1': EXPECTED[0][2],
                          'A2': EXPECTED[1][0], 'C2': EXPECTED[1][2]}
        for ref, value in expected_cells.items():
            cell = cells[ref]
            require(cell.attrib.get('t') == 'inlineStr', f'{ref} not explicitly text')
            require(cell.find('s:f', NS) is None, f'{ref} became a formula')
            text = cell.find('s:is/s:t', NS)
            require(text is not None and text.text == value, f'{ref} XML text changed')
        require('B2' not in cells, 'empty field must clear B2')
        require(cells['D3'].find('s:f', NS) is not None, 'unrelated formula missing')
        require(cells['D3'].find('s:f', NS).text == '3+4', 'unrelated formula changed')
    book = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    sheet = book.active
    for ref, value in expected_cells.items():
        require(sheet[ref].data_type == 's' and sheet[ref].value == value, f'{ref} reader/type mismatch')
    require(sheet['B2'].value is None, 'reader reports cleared cell populated')
    require(sheet['D3'].data_type == 'f' and sheet['D3'].value == '=3+4', 'reader lost unrelated formula')
    cached = load_workbook(io.BytesIO(data), data_only=True, read_only=False).active
    require(cached['D3'].value == 7, 'formula cached value mismatch (not independent recalculation)')
    return checks


def main():
    n = validate()
    print(f'R58 independent TSV/ZIP/XML/openpyxl checks: {n} predicates passed; not Excel or LibreOffice rendering/recalculation')


if __name__ == '__main__':
    main()
