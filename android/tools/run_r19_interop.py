#!/usr/bin/env python3
"""Independent Python task-frame/state references and openpyxl consumption of actual Kotlin XLSX bytes."""
import base64,datetime,hashlib,io,json,os,random,struct,subprocess,uuid,zlib,zipfile
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'testing/build/r19';rng=random.Random(1903)
for line in (OUT/'result.txt').read_text().splitlines():
    if len(line)>66 and line[64:66]=='  ':assert hashlib.sha256((ROOT/line[66:]).read_bytes()).hexdigest()==line[:64],line
b64=lambda b:base64.b64encode(b).decode()
epoch=datetime.date(1970,1,1)
def frame(rev,tasks):
    raw=struct.pack('>4sqI',b'FTS2',rev,len(tasks))
    for t in tasks:
        title=t['title'].encode();day=(t['date']-epoch).days if t['date'] else -(2**63)
        raw+=t['id'].encode()+struct.pack('>BH',int(t['done'])+2*int(t['archived']),len(title))+title+struct.pack('>q',day)
    return b64(raw+struct.pack('>I',zlib.crc32(raw)))
requests=[];expected=[]
for i in range(150):
    tasks=[dict(id=str(uuid.UUID(int=rng.getrandbits(128))),title=f'Task {i} 界🙂 {j}',date=None if j%3==0 else datetime.date(2028,2,29),done=bool(j%2),archived=j>2) for j in range(5)]
    if i<50:
        requests.append('T\t'+frame(i,tasks));expected.append(('task_report',tasks))
    incoming=tasks[:2]+[dict(id=str(uuid.UUID(int=rng.getrandbits(128))),title='Imported '+str(j),date=None,done=True,archived=bool(j%2)) for j in range(i%4)]
    added=incoming[2:];after=frame(i+1,tasks+added) if added else frame(i,tasks)
    undo=frame(i+2,tasks) if added else '-'
    requests.append('M\t'+frame(i,tasks)+'\t'+frame(100000,incoming));expected.append(('literal',f'{len(added)}:2:{after}:{undo}'))
    title='Edited α '+str(i);date=datetime.date(9999,12,31) if i%2 else None
    revised=[tasks[0]|{'title':title,'date':date}]+tasks[1:]
    requests.append('\t'.join(['U',frame(i,tasks),b64(title.encode()),str(date) if date else '']))
    expected.append(('literal',frame(i+1,revised)+':'+frame(i+2,tasks)))
for i in range(80):
    count=i%16+1;text='=1+1 '+''.join(rng.choice(['a','界','🙂','&','<','>','\n','\r','\t',"'"]) for _ in range(50));number=str(i+0.25)
    requests.append('\t'.join(['W',str(count),b64(text.encode()),number]));expected.append(('workbook',(count,text,float(number))))
java=str(Path(os.environ['JAVA_HOME'])/'bin/java') if 'JAVA_HOME' in os.environ else 'java'
cp=os.pathsep.join(str(OUT/f) for f in ['checks.jar','TaskChanges.jar','LiteralWorkbook.jar'])
r=subprocess.run([java,'-cp',cp,'dev.fusion.testing.R19Oracle'],input='\n'.join(requests)+'\n',capture_output=True,text=True,check=True)
actual=r.stdout.splitlines();assert len(actual)==len(expected),(len(actual),len(expected))
for a,(kind,value) in zip(actual,expected):
    if kind=='literal':assert a==value,(a,value)
    elif kind=='task_report':
        book=openpyxl.load_workbook(io.BytesIO(base64.b64decode(a)),data_only=False)
        assert book.sheetnames==['Active','Archived']
        for archived,name in [(False,'Active'),(True,'Archived')]:
            sheet=book[name];rows=list(sheet.values);assert rows[0]==('ID','Title','Date (text)','Completed','Archived')
            expected_tasks=[t for t in value if t['archived']==archived];assert len(rows)==len(expected_tasks)+1
            for row,t in zip(rows[1:],expected_tasks):
                assert row[:2]==(t['id'],t['title'])
                assert row[2]==str(t['date']) if t['date'] else row[2] in (None,'')
                assert row[3:]==(t['done'],t['archived'])
    else:
        count,text,number=value;data=base64.b64decode(a);book=openpyxl.load_workbook(io.BytesIO(data),data_only=False)
        assert book.sheetnames==[f'Sheet{i}' for i in range(1,count+1)]
        for i,sheet in enumerate(book.worksheets,1):
            assert sheet['A1'].value==text and sheet['A1'].data_type=='s'
            assert sheet['B1'].value==number and sheet['C1'].value==(i%2==0)
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            assert len(z.namelist())==count+4
            assert not any('vba' in name.lower() for name in z.namelist())
            assert not any(b'<f>' in z.read(name) for name in z.namelist())
receipt={'result':'PASS','cases':len(expected),'breakdown':{'merge_reference':150,'edit_undo_reference':150,'multi_sheet_openpyxl_readback':80,'task_report_openpyxl_readback':50},'openpyxl':openpyxl.__version__,
         'scope':'Actual separately compiled Kotlin outputs; independent Python struct/zlib and external openpyxl. Same-author corpus, not Excel/LibreOffice fidelity or expert review.',
         'sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [Path(__file__),OUT/'checks.jar',OUT/'TaskChanges.jar',OUT/'LiteralWorkbook.jar']}}
(ROOT/'docs/verification/r19-interop.json').write_text(json.dumps(receipt,indent=2)+'\n')
print(f'R19: {len(expected)} independent/reference cases PASS, including130 workbook/report openpyxl readbacks')
